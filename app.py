import os
import sys
import shutil
from pathlib import Path
from typing import Optional
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.requests import Request
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel
import openpyxl
import tempfile

from database import init_db, get_connection, BASE_DIR

app = FastAPI(title="Applicants System")

UPLOADS_DIR = os.path.join(BASE_DIR, "uploads")
PHOTOS_DIR = os.path.join(UPLOADS_DIR, "photos")
DOCUMENTS_DIR = os.path.join(UPLOADS_DIR, "documents")

os.makedirs(PHOTOS_DIR, exist_ok=True)
os.makedirs(DOCUMENTS_DIR, exist_ok=True)

STATIC_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static")
TEMPLATES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates")

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
app.mount("/uploads", StaticFiles(directory=UPLOADS_DIR), name="uploads")

templates = Jinja2Templates(directory=TEMPLATES_DIR)

init_db()


# ─── Page Routes ─────────────────────────────────────────────────────────────

@app.get("/")
async def index(request: Request):
    return templates.TemplateResponse(request, "index.html")

@app.get("/applicant/{id}")
async def applicant_page(request: Request, id: int):
    return templates.TemplateResponse(request, "applicant.html", {"applicant_id": id})

@app.get("/edit/{id}")
async def edit_page(request: Request, id: int):
    return templates.TemplateResponse(request, "edit.html", {"applicant_id": id})

@app.get("/documents/{id}")
async def documents_page(request: Request, id: int):
    return templates.TemplateResponse(request, "documents.html", {"applicant_id": id})


# ─── API: Specializations ─────────────────────────────────────────────────────

@app.get("/api/specializations")
async def get_specializations():
    conn = get_connection()
    rows = conn.execute("SELECT code, name FROM specializations ORDER BY code").fetchall()
    conn.close()
    return [dict(r) for r in rows]


# ─── API: Applicants ──────────────────────────────────────────────────────────

@app.get("/api/applicants")
async def get_applicants():
    conn = get_connection()
    rows = conn.execute("""
        SELECT id, last_name_ru, first_name_ru, last_name_cn, first_name_cn,
               specialization_code, specialization_name, email, password, snils,
               passport_number, passport_expiry, passport_issued_by,
               education_type, doc1_number, doc2_number, partner, photo_path
        FROM applicants ORDER BY id
    """).fetchall()
    conn.close()
    return [dict(r) for r in rows]


@app.post("/api/applicants")
async def create_applicant():
    conn = get_connection()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO applicants (last_name_ru, first_name_ru, last_name_cn, first_name_cn,
            email, password, snils, passport_number, passport_expiry,
            passport_issued_by, education_type, specialization_code,
            specialization_name, doc1_number, doc2_number, partner, photo_path)
        VALUES ('', '', '', '', '', '', '', '', '', '', 'после школы', '', '', '', '', '', NULL)
    """)
    new_id = cur.lastrowid
    conn.commit()
    conn.close()
    return {"id": new_id}


@app.get("/api/applicants/{id}")
async def get_applicant(id: int):
    conn = get_connection()
    row = conn.execute("SELECT * FROM applicants WHERE id = ?", (id,)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(status_code=404, detail="Абитуриент не найден")
    return dict(row)


class ApplicantUpdate(BaseModel):
    last_name_ru: Optional[str] = None
    first_name_ru: Optional[str] = None
    last_name_cn: Optional[str] = None
    first_name_cn: Optional[str] = None
    email: Optional[str] = None
    password: Optional[str] = None
    snils: Optional[str] = None
    passport_number: Optional[str] = None
    passport_expiry: Optional[str] = None
    passport_issued_by: Optional[str] = None
    education_type: Optional[str] = None
    specialization_code: Optional[str] = None
    specialization_name: Optional[str] = None
    doc1_number: Optional[str] = None
    doc2_number: Optional[str] = None
    partner: Optional[str] = None


@app.put("/api/applicants/{id}")
async def update_applicant(id: int, data: ApplicantUpdate):
    conn = get_connection()
    existing = conn.execute("SELECT id FROM applicants WHERE id = ?", (id,)).fetchone()
    if not existing:
        conn.close()
        raise HTTPException(status_code=404, detail="Абитуриент не найден")

    fields = {k: v for k, v in data.dict().items() if v is not None}
    if fields:
        set_clause = ", ".join(f"{k} = ?" for k in fields)
        values = list(fields.values()) + [id]
        conn.execute(f"UPDATE applicants SET {set_clause} WHERE id = ?", values)
        conn.commit()
    conn.close()
    return {"ok": True}


@app.delete("/api/applicants/{id}")
async def delete_applicant(id: int):
    conn = get_connection()
    conn.execute("DELETE FROM applicants WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    return {"ok": True}


# ─── API: Export Excel ────────────────────────────────────────────────────────

@app.get("/api/export/applicants")
async def export_applicants():
    conn = get_connection()
    rows = conn.execute("SELECT * FROM applicants ORDER BY id").fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Абитуриенты"

    headers = [
        "ID", "Фамилия (рус)", "Имя (рус)", "Фамилия (кит)", "Имя (кит)",
        "Email", "Пароль", "СНИЛС", "Номер паспорта", "Срок паспорта",
        "Кем выдан паспорт", "Тип образования", "Код специализации",
        "Название специализации", "Номер документа 1", "Номер документа 2", "Партнёр"
    ]
    ws.append(headers)

    for row in rows:
        d = dict(row)
        ws.append([
            d.get("id"), d.get("last_name_ru"), d.get("first_name_ru"),
            d.get("last_name_cn"), d.get("first_name_cn"), d.get("email"),
            d.get("password"), d.get("snils"), d.get("passport_number"),
            d.get("passport_expiry"), d.get("passport_issued_by"),
            d.get("education_type"), d.get("specialization_code"),
            d.get("specialization_name"), d.get("doc1_number"),
            d.get("doc2_number"), d.get("partner")
        ])

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(tmp.name)
    tmp.close()

    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="applicants.xlsx"
    )


# ─── API: Photo Upload ────────────────────────────────────────────────────────

@app.post("/api/upload/photo")
async def upload_photo(applicant_id: int = Form(...), file: UploadFile = File(...)):
    ext = Path(file.filename).suffix
    filename = f"photo_{applicant_id}{ext}"
    dest = os.path.join(PHOTOS_DIR, filename)

    with open(dest, "wb") as f:
        shutil.copyfileobj(file.file, f)

    rel_path = f"photos/{filename}"
    conn = get_connection()
    conn.execute("UPDATE applicants SET photo_path = ? WHERE id = ?", (rel_path, applicant_id))
    conn.commit()
    conn.close()

    return {"path": rel_path, "url": f"/uploads/{rel_path}"}


# ─── API: Documents ───────────────────────────────────────────────────────────

@app.get("/api/documents/{id}/{doc_type}")
async def get_document(id: int, doc_type: str):
    conn = get_connection()
    rows = conn.execute(
        "SELECT * FROM documents WHERE applicant_id = ? AND doc_type = ?",
        (id, doc_type)
    ).fetchall()
    conn.close()
    if not rows:
        return []
    return [dict(r) for r in rows]


class DocumentUpdate(BaseModel):
    lang: Optional[str] = "ru"
    number: Optional[str] = None
    issued_date: Optional[str] = None
    issued_by: Optional[str] = None
    issue_place: Optional[str] = None
    birth_province: Optional[str] = None
    institution_name: Optional[str] = None
    issue_date: Optional[str] = None


@app.put("/api/documents/{id}/{doc_type}")
async def update_document(id: int, doc_type: str, data: DocumentUpdate):
    conn = get_connection()
    existing = conn.execute(
        "SELECT id FROM documents WHERE applicant_id = ? AND doc_type = ? AND lang = ?",
        (id, doc_type, data.lang)
    ).fetchone()

    fields = {k: v for k, v in data.dict().items() if v is not None and k != "lang"}

    if existing:
        if fields:
            set_clause = ", ".join(f"{k} = ?" for k in fields)
            values = list(fields.values()) + [id, doc_type, data.lang]
            conn.execute(
                f"UPDATE documents SET {set_clause} WHERE applicant_id = ? AND doc_type = ? AND lang = ?",
                values
            )
    else:
        cols = ["applicant_id", "doc_type", "lang"] + list(fields.keys())
        vals = [id, doc_type, data.lang] + list(fields.values())
        placeholders = ", ".join(["?"] * len(vals))
        conn.execute(f"INSERT INTO documents ({', '.join(cols)}) VALUES ({placeholders})", vals)

    conn.commit()
    conn.close()
    return {"ok": True}


@app.post("/api/upload/document")
async def upload_document(
    applicant_id: int = Form(...),
    doc_type: str = Form(...),
    lang: str = Form("ru"),
    file: UploadFile = File(...)
):
    ext = Path(file.filename).suffix
    filename = f"doc_{applicant_id}_{doc_type}_{lang}{ext}"
    dest = os.path.join(DOCUMENTS_DIR, filename)

    with open(dest, "wb") as f:
        shutil.copyfileobj(file.file, f)

    rel_path = f"documents/{filename}"
    conn = get_connection()
    existing = conn.execute(
        "SELECT id FROM documents WHERE applicant_id = ? AND doc_type = ? AND lang = ?",
        (applicant_id, doc_type, lang)
    ).fetchone()

    if existing:
        conn.execute(
            "UPDATE documents SET scan_path = ? WHERE applicant_id = ? AND doc_type = ? AND lang = ?",
            (rel_path, applicant_id, doc_type, lang)
        )
    else:
        conn.execute(
            "INSERT INTO documents (applicant_id, doc_type, lang, scan_path) VALUES (?, ?, ?, ?)",
            (applicant_id, doc_type, lang, rel_path)
        )

    conn.commit()
    conn.close()
    return {"path": rel_path, "url": f"/uploads/{rel_path}"}
